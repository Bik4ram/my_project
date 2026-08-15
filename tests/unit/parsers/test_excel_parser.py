from datetime import datetime, timezone
from pathlib import Path

import openpyxl

from semrag.ingestion.parsers.documents.excel_parser import ExcelParser
from semrag.ingestion.scanner.models import FileCategory, FileRecord


def _record(path: Path) -> FileRecord:
    return FileRecord(
        project_id="proj1",
        absolute_path=str(path),
        relative_path=path.name,
        size_bytes=path.stat().st_size,
        mtime=datetime.now(timezone.utc),
        content_hash="deadbeef",
        category=FileCategory.EXCEL,
        extension=".xlsx",
        is_binary=True,
    )


def test_excel_extracts_headers_and_sample_rows(tmp_path: Path):
    path = tmp_path / "register_map.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registers"
    ws.append(["Address", "Register Name", "Width"])
    ws.append(["0x00", "CTRL", 8])
    ws.append(["0x04", "STATUS", 8])
    ws.append(["0x08", "DATA", 32])
    wb.save(path)

    result = ExcelParser().parse(path, _record(path))
    assert result.success is True

    sheets = result.payload["sheets"]
    assert len(sheets) == 1
    sheet = sheets[0]
    assert sheet["name"] == "Registers"
    assert sheet["headers"] == ["Address", "Register Name", "Width"]
    assert sheet["row_count"] == 3
    assert sheet["sample_rows"][0][1] == "CTRL"


def test_excel_multiple_sheets(tmp_path: Path):
    path = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["a", "b"])
    ws1.append([1, 2])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["x", "y", "z"])
    ws2.append([1, 2, 3])
    wb.save(path)

    result = ExcelParser().parse(path, _record(path))
    sheet_names = {s["name"] for s in result.payload["sheets"]}
    assert sheet_names == {"Sheet1", "Sheet2"}


def test_excel_invalid_file_fails_gracefully(tmp_path: Path):
    fake = tmp_path / "not_really.xlsx"
    fake.write_bytes(b"not an excel file")

    result = ExcelParser().parse(fake, _record(fake))
    assert result.success is False
    assert result.has_errors is True
